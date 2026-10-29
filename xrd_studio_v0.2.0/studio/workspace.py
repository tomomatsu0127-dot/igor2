"""Igor-independent data, project, snapshot-figure and recovery operations."""
from __future__ import annotations

import copy
import csv
import json
from pathlib import Path
import re

import numpy as np
from scipy.signal import savgol_filter

from . import analysis, io
from .document import uid, save_document, new_state


def status(sample):
    reviewed = {a['peak_id'] for a in sample['assignments']}
    n = len(sample['peaks'])
    return '未解析' if not n else ('同定完了' if all(p['id'] in reviewed for p in sample['peaks']) else '同定中')


def touch(state, sample_id):
    sample = state['samples'][sample_id]
    sample['revision'] = sample.get('revision', 0) + 1
    sample['status'] = status(sample)
    for plot in state['plots'].values():
        if sample_id in plot['source_ids']:
            plot['stale'] = True


def new_project(state, name):
    name = str(name).strip()
    if not name or any(p['name'] == name and not p.get('deleted') for p in state['projects'].values()):
        raise ValueError('空でない、重複しないプロジェクト名を指定してください。')
    pid = uid()
    state['projects'][pid] = {'id':pid, 'name':name, 'sample_ids':[], 'compare_ids':[], 'deleted':False}
    state['active_project'] = pid
    return pid


def set_metadata(state, sid, serial, name, composition):
    s = state['samples'][sid]
    serial = str(serial).strip()
    if not serial or not str(name).strip():
        raise ValueError('通し番号と表示名を入力してください。')
    if any(x != sid and state['samples'][x]['serial'] == serial and not state['samples'][x].get('deleted') for x in state['projects'][s['project_id']]['sample_ids']):
        raise ValueError('同じプロジェクトに同じ通し番号があります。')
    s.update(serial=serial, name=str(name).strip(), composition=str(composition))


def add_measurement(state, path, serial, name, composition):
    sample = io.make_sample(path, state['active_project'], serial, name, composition)
    sid = sample['id']
    state['samples'][sid] = sample
    state['projects'][sample['project_id']]['sample_ids'].append(sid)
    set_metadata(state, sid, serial, name, composition)
    sample.update(deleted=False, revision=0)
    return sid


def copy_sample(state, sid, pid, serial):
    source = state['samples'][sid]
    target = copy.deepcopy(source)
    target.update(id=uid(), project_id=pid, serial=str(serial), copied_from=sid, copied_at=analysis.timestamp(), deleted=False)
    state['samples'][target['id']] = target
    state['projects'][pid]['sample_ids'].append(target['id'])
    set_metadata(state, target['id'], serial, target['name'], target['composition'])
    return target['id']


def delete_sample(state, sid, deleted=True):
    s = state['samples'][sid]
    if not deleted:
        set_metadata(state, sid, s['serial'], s['name'], s['composition'])
    s['deleted'] = deleted
    if deleted:
        for project in state['projects'].values():
            project['compare_ids'] = [x for x in project['compare_ids'] if x != sid]


def add_reference(state, path):
    ref = io.read_reference(path)
    for existing in state['references'].values():
        if existing['scientific_hash'] == ref['scientific_hash']:
            return existing['id']
    state['references'][ref['id']] = ref
    return ref['id']


def new_wave(state, sid, name=None):
    sample = state['samples'][sid]
    wave_id = uid()
    state['waves'][wave_id] = {'id':wave_id, 'sample_id':sid, 'name':name or sample['serial']+' の作業データ',
        'x':sample['raw_x'].copy(), 'y':sample['processed_y'].copy(), 'x_unit':'deg', 'y_unit':'%',
        'note':'規格化データの独立した作業コピー。解析結果へ自動反映しません。',
        'source':'processed_y', 'operations':[]}
    return wave_id


def set_wave_value(state, wid, row, value):
    value = float(value)
    wave = state['waves'][wid]
    if not np.isfinite(value) or not 0 <= row < len(wave['y']):
        raise ValueError('表には有限の数値を入力してください。')
    before = float(wave['y'][row])
    wave['y'][row] = value
    wave['operations'].append({'kind':'cell', 'row':row, 'before':before, 'value':value})


def transform_wave(state, wid, kind, value=1):
    wave = state['waves'][wid]
    value = float(value)
    if not np.isfinite(value):
        raise ValueError('有限の数値を指定してください。')
    y, x = wave['y'], wave['x']
    if kind == 'scale':
        result = y * value
    elif kind == 'offset':
        result = y + value
    elif kind == 'normalize':
        result = io.normalize(y)
        wave['y_unit'] = '%'
    elif kind == 'derivative':
        if len(x) < 3 or not np.all(np.diff(x) > 0): raise ValueError('微分には昇順の3点以上が必要です。')
        result = np.gradient(y, x)
        wave['y_unit'] += '/'+wave['x_unit']
    elif kind == 'smooth':
        width = int(value)
        if width < 3 or width % 2 == 0 or width > len(y): raise ValueError('平滑化幅は点数以下の3以上の奇数です。')
        result = savgol_filter(y, width, min(2,width-1))
    else:
        raise ValueError('不明な処理です。')
    if not np.all(np.isfinite(result)): raise ValueError('処理後の値が有限でありません。')
    wave['y'] = result
    wave['operations'].append({'kind':kind,'value':value,'created':analysis.timestamp()})


def external_peaks(sample, path):
    """Validate the complete list before mutation and retain entered angles exactly."""
    parsed = []
    for line_no, line in enumerate(io.read_text(path).splitlines(), 1):
        line = line.strip()
        if not line or line.startswith(('#','//')): continue
        fields = io.tokens(line)
        try:
            if len(fields) not in (1,2): raise ValueError('角度、または角度と強度の1-2列が必要です。')
            angle = io.finite_number(fields[0])
            height = io.finite_number(fields[1]) if len(fields)==2 else None
        except ValueError:
            if not parsed and re.match(r'^(2theta|angle|tt)\b',line,re.I):continue
            raise ValueError(f'{line_no}行目のピーク値を確認してください。')
        if not sample['raw_x'][0] <= angle <= sample['raw_x'][-1]:raise ValueError(f'{line_no}行目の角度が測定範囲外です。')
        if any(abs(angle-x[0])<1e-10 for x in parsed):raise ValueError('外部ピークの角度が重複しています。')
        parsed.append((angle,height))
    if not parsed:raise ValueError('外部ピークがありません。')
    for angle,height in parsed:
        if any(abs(angle-p['angle'])<1e-10 for p in sample['peaks']):continue
        index=int(np.argmin(abs(sample['raw_x']-angle)))
        sample['peaks'].append({'id':sample['next_peak_id'],'index':index,'angle':angle,
            'height':float(sample['processed_y'][index]) if height is None else height,
            'origin':'External','reviewed':False,'source_name':Path(path).name})
        sample['next_peak_id']+=1
    sample['peaks'].sort(key=lambda p:p['angle'])
    analysis.invalidate_lattice(sample)


def evidence(sample, reference, shift, tolerance):
    explained={a['peak_id'] for a in sample['assignments'] if a['ref_id'] is not None}
    result=[]
    for i,peak,delta in analysis._matches(sample,reference,shift,tolerance):
        result.append([i, float(reference['x'][i]+shift), float(reference['intensity'][i]),
            peak['id'] if peak else None, delta,
            '他相採用済みと一致' if peak and peak['id'] in explained else ('一致' if peak else '不足')])
    return sorted(result,key=lambda row:row[1])


def save_plot(state, name, source_ids):
    if not source_ids:raise ValueError('比較する試料をチェックしてください。')
    snapshot=new_state()
    pid=snapshot['active_project']
    snapshot['projects'][pid]['compare_ids']=list(source_ids)
    snapshot['projects'][pid]['sample_ids']=list(source_ids)
    snapshot['references']=copy.deepcopy(state['references'])
    snapshot['plot']=copy.deepcopy(state['plot'])
    for sid in source_ids:
        sample=copy.deepcopy(state['samples'][sid]);sample['project_id']=pid
        snapshot['samples'][sid]=sample
    plot_id=uid()
    state['plots'][plot_id]={'id':plot_id,'name':name,'created':analysis.timestamp(),
        'source_ids':list(source_ids),'snapshot':snapshot,'stale':False}
    return plot_id


def recover(document, directory):
    directory=Path(directory)
    directory.mkdir(parents=True,exist_ok=True)
    stem=document.state['id']
    newest=directory/(stem+'-1.labx')
    # Write successfully before rotating. Failed saves preserve every old generation.
    pending=directory/(stem+'-pending.labx')
    save_document(document,pending,mark_saved=False)
    for older in (2,1):
        source=directory/f'{stem}-{older}.labx'
        if source.exists():source.replace(directory/f'{stem}-{older+1}.labx')
    pending.replace(newest)
    return newest


def export_tables(state, directory, xlsx=False):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    tables={
        'Samples':[['sample_id','project','serial','name','composition','status','deleted','source_name','source_sha256']],
        'Peaks':[['sample_id','peak_id','2theta_deg','height_percent','origin','assignment']],
        'Assignments':[['sample_id','peak_id','reference_id','reference_revision','reflection','h','k','l','shift_deg','mode']],
        'LatticeRuns':[['sample_id','run_id','crystal','a_A','a_sigma_A','c_A','c_sigma_A','wavelength_A','zero_deg','stale']],
        'Settings':[['sample_id','parameter','value']],
        'References':[['reference_id','revision','card','name','formula','phase_class','scientific_hash']]}
    for sid,s in state['samples'].items():
        tables['Samples'].append([sid,state['projects'][s['project_id']]['name'],s['serial'],s['name'],s['composition'],status(s),s.get('deleted',False),s['source_name'],s['source_sha256']])
        for p in s['peaks']:
            tables['Peaks'].append([sid,p['id'],p['angle'],p['height'],p['origin'],analysis_status_label(s,p['id'],state['references'])])
        for a in s['assignments']:
            ref=state['references'].get(a['ref_id']);i=a['reflection']
            tables['Assignments'].append([sid,a['peak_id'],a['ref_id'],ref['revision'] if ref else None,i,*[float(ref[k][i]) if ref and np.isfinite(ref[k][i]) else None for k in ('h','k','l')],a['shift'],a['mode']])
        for r in s['lattice_runs']:
            tables['LatticeRuns'].append([sid,r['id'],r['crystal'],*[r[k] for k in ('a_A','a_sigma_A','c_A','c_sigma_A','wavelength_A','zero_deg','stale')]])
        for k,v in s['parameters'].items():tables['Settings'].append([sid,k,v])
    for rid,r in state['references'].items():tables['References'].append([rid,r['revision'],r['card'],r['name'],r['formula'],r['phase_class'],r['scientific_hash']])
    for name,rows in tables.items():
        with (directory/(name+'.csv')).open('w',newline='',encoding='utf-8-sig') as stream:csv.writer(stream).writerows(rows)
    if xlsx:
        from openpyxl import Workbook
        book=Workbook();book.remove(book.active)
        for name,rows in tables.items():
            sheet=book.create_sheet(name)
            for row in rows:
                sheet.append(row)
                # User labels are text, never spreadsheet formulas.
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value,str):cell.data_type='s'
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=22
        book.save(directory/'XRD_results.xlsx')
    return tables


def analysis_status_label(sample,pid,references):
    from .plotting import assignment_label
    return assignment_label(sample,pid,references)
