import copy
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import numpy as np
from studio import analysis,demo,workspace
from studio.document import Document,save_document,load_document,validate,uid


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)
        self.doc=Document();demo.populate(self.doc.state);self.sid=next(iter(self.doc.state['samples']));self.sample=self.doc.state['samples'][self.sid]
    def detect(self):analysis.detect_peaks(self.sample,self.sample['parameters'])
    def test_working_copy_cell_undo_and_restart_preserve_raw(self):
        raw=self.sample['raw_y'].copy();wid=self.doc.change('copy',lambda st:workspace.new_wave(st,self.sid))
        old=float(self.doc.state['waves'][wid]['y'][0]);self.doc.change('edit',lambda st:workspace.set_wave_value(st,wid,0,321.5))
        np.testing.assert_array_equal(self.doc.state['samples'][self.sid]['raw_y'],raw)
        self.doc.undo();self.assertEqual(self.doc.state['waves'][wid]['y'][0],old);self.doc.redo()
        path=self.root/'wave.labx';save_document(self.doc,path);d=load_document(path)
        self.assertEqual(d.state['waves'][wid]['y'][0],321.5);self.assertTrue(d.undo());self.assertEqual(d.state['waves'][wid]['y'][0],old)
    def test_copy_transforms_and_units_without_mutating_source(self):
        wid=workspace.new_wave(self.doc.state,self.sid);w=self.doc.state['waves'][wid];original=self.sample['processed_y'].copy()
        workspace.transform_wave(self.doc.state,wid,'scale',2);workspace.transform_wave(self.doc.state,wid,'offset',3)
        np.testing.assert_allclose(w['y'],2*original+3);workspace.transform_wave(self.doc.state,wid,'derivative')
        self.assertEqual(w['y_unit'],'%/deg');np.testing.assert_array_equal(self.sample['processed_y'],original)
        with self.assertRaises(ValueError):workspace.transform_wave(self.doc.state,wid,'smooth',4)
    def test_invalid_cell_rolls_back_history_and_values(self):
        wid=workspace.new_wave(self.doc.state,self.sid);before=self.doc.state['waves'][wid]['y'].copy()
        with self.assertRaises(ValueError):self.doc.change('invalid',lambda st:workspace.set_wave_value(st,wid,0,float('nan')))
        self.assertEqual(len(self.doc.undo_stack),0);np.testing.assert_array_equal(self.doc.state['waves'][wid]['y'],before)
    def test_cross_project_comparison_and_copy_keep_identification(self):
        self.detect();p=self.sample['peaks'][0];analysis.assign(self.sample,p['id'])
        pid=workspace.new_project(self.doc.state,'別プロジェクト');sid2=workspace.copy_sample(self.doc.state,self.sid,pid,'015')
        self.doc.state['projects'][pid]['compare_ids']=[self.sid,sid2];validate(self.doc.state)
        path=self.root/'cross.labx';save_document(self.doc,path);d=load_document(path)
        self.assertEqual(d.state['projects'][pid]['compare_ids'],[self.sid,sid2]);self.assertEqual(d.state['samples'][sid2]['assignments'],self.sample['assignments'])
        self.assertEqual(d.state['samples'][sid2]['copied_from'],self.sid)
    def test_serial_leading_zero_duplicate_and_restore_conflict(self):
        workspace.set_metadata(self.doc.state,self.sid,'015','表示名','Fe2O3');workspace.delete_sample(self.doc.state,self.sid)
        other=list(self.doc.state['samples'])[1];workspace.set_metadata(self.doc.state,other,'015','別試料','')
        with self.assertRaises(ValueError):self.doc.change('restore',lambda st:workspace.delete_sample(st,self.sid,False))
        self.assertTrue(self.doc.state['samples'][self.sid]['deleted']);self.assertEqual(self.doc.state['samples'][self.sid]['serial'],'015')
    def test_saved_plot_snapshot_survives_analysis_and_trash(self):
        state=self.doc.state;plotid=workspace.save_plot(state,'保存図',[self.sid]);before=state['plots'][plotid]['snapshot']['samples'][self.sid]['processed_y'].copy()
        analysis.add_peak(self.sample,100);workspace.touch(state,self.sid);workspace.delete_sample(state,self.sid)
        self.assertTrue(state['plots'][plotid]['stale']);np.testing.assert_array_equal(state['plots'][plotid]['snapshot']['samples'][self.sid]['processed_y'],before)
        save_document(self.doc,self.root/'plot.labx');d=load_document(self.root/'plot.labx');self.assertIn(plotid,d.state['plots'])
    def test_external_angles_exact_and_atomic_rejection(self):
        path=self.root/'peaks.csv';path.write_text('angle,height\n30.1234567,55\n31.3333333,12\n',encoding='utf-8')
        workspace.external_peaks(self.sample,path);self.assertEqual(self.sample['peaks'][0]['angle'],30.1234567)
        before=copy.deepcopy(self.sample['peaks']);path.write_text('32,30\n400,30\n',encoding='utf-8')
        with self.assertRaises(ValueError):workspace.external_peaks(self.sample,path)
        self.assertEqual(self.sample['peaks'],before)
    def test_external_is_retained_after_detection(self):
        path=self.root/'p.txt';path.write_text('30.1234567\n');workspace.external_peaks(self.sample,path);pid=self.sample['peaks'][0]['id'];self.detect()
        self.assertEqual(next(p for p in self.sample['peaks'] if p['id']==pid)['angle'],30.1234567)
    def test_element_set_filter_and_protected_types(self):
        r={'formula':'Fe2O3','phase_class':''};self.assertTrue(analysis.formula_compatible(r,'Fe O'));self.assertFalse(analysis.formula_compatible(r,'F O'))
        r['phase_class']='M-type';self.assertTrue(analysis.formula_compatible(r,'Ba O'))
        r={'formula':'','phase_class':''};self.assertTrue(analysis.formula_compatible(r,'Ba O'))
    def test_candidate_modes_and_new_phase_score(self):
        self.detect();p=self.sample['peaks'][0];refs=self.doc.state['references'];c=analysis.score_candidates(self.sample,refs,p['id'],self.sample['parameters'])[0]
        self.assertEqual(analysis.score_candidates(self.sample,refs,p['id'],self.sample['parameters'],candidate_mode='adopted'),[])
        n=analysis.score_candidates(self.sample,refs,p['id'],self.sample['parameters'],candidate_mode='new')[0]
        expected=np.clip(.3*n['local']+.4*n['global']+.2*n['new']+.1*n['intensity']-min(15,5*n['missing']),0,100)
        self.assertAlmostEqual(n['total'],expected);analysis.assign(self.sample,p['id'],c)
        self.assertEqual(analysis.score_candidates(self.sample,refs,p['id'],self.sample['parameters'],candidate_mode='new'),[])
    def test_background_off_and_explicit_adoption(self):
        self.detect();p=self.sample['peaks'][0];angle=p['angle'];r=analysis.fit_profile(self.sample,p['id'],angle-.4,angle+.4,background=False)
        np.testing.assert_array_equal(r['baseline'],0);self.assertEqual(p['angle'],angle)
    def test_recovery_rotation_and_failed_write_preserves_generation(self):
        for i in range(4):self.doc.state['plot']['title']=str(i);workspace.recover(self.doc,self.root)
        files=list(self.root.glob('*.labx'));self.assertEqual(len(files),3)
        newest=self.root/(self.doc.state['id']+'-1.labx');before=newest.read_bytes()
        with patch('studio.document.os.replace',side_effect=OSError('disk failure')):
            with self.assertRaises(OSError):workspace.recover(self.doc,self.root)
        self.assertEqual(newest.read_bytes(),before)
    def test_xlsx_and_csv_keep_serial_and_text_not_formulas(self):
        self.sample.update(serial='015',name='=1+1');workspace.export_tables(self.doc.state,self.root,xlsx=True)
        from openpyxl import load_workbook
        book=load_workbook(self.root/'XRD_results.xlsx');sheet=book['Samples']
        self.assertEqual(sheet['C2'].value,'015');self.assertEqual(sheet['D2'].data_type,'s');self.assertEqual(sheet['D2'].value,'=1+1')
    def test_complete_unknown_is_distinct_from_unreviewed(self):
        self.detect();self.assertEqual(workspace.status(self.sample),'同定中')
        for p in self.sample['peaks']:analysis.assign(self.sample,p['id'])
        self.assertEqual(workspace.status(self.sample),'同定完了')
        analysis.clear_assignments(self.sample,self.sample['peaks'][0]['id']);self.assertEqual(workspace.status(self.sample),'同定中')


if __name__=='__main__':unittest.main()
